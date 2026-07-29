from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUTPUT_PATH = Path(r"C:\Users\hp\Documents\StudySync_SAD_Report.docx")

BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"
DARK_GRAY = "404040"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size=11, bold=False, color=None, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_text(paragraph, text, size=11, bold=False, color=None, italic=False):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color, italic=italic)
    return run


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    add_text(paragraph, text, size={1: 14, 2: 12, 3: 11}.get(level, 11), bold=True, color=BLUE)
    return paragraph


def add_body(doc, text, bold_prefix=None):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix and text.startswith(bold_prefix):
        add_text(paragraph, bold_prefix, bold=True)
        add_text(paragraph, text[len(bold_prefix):])
    else:
        add_text(paragraph, text)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    add_text(paragraph, text)
    return paragraph


def add_numbered(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(3)
    add_text(paragraph, text)
    return paragraph


def set_cell_text(cell, text, bold=False, color=None, size=9.5, align=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    for index, line in enumerate(str(text).split("\n")):
        if index:
            paragraph.add_run().add_break()
        add_text(paragraph, line, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)


def add_table(doc, headers, rows, widths=None, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, BLUE)
        set_cell_text(
            cell,
            header,
            bold=True,
            color=WHITE,
            size=font_size,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if row_index % 2 == 1:
                set_cell_shading(cells[index], LIGHT_GRAY)
            set_cell_text(cells[index], value, size=font_size)

    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)

    doc.add_paragraph()
    return table


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    add_text(paragraph, text, size=10, bold=True, italic=True)
    return paragraph


def add_diagram(doc, title, rows, column_widths=None):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for row_index, values in enumerate(rows):
        for col_index, value in enumerate(values):
            cell = table.rows[row_index].cells[col_index]
            if value:
                fill = LIGHT_BLUE if value.startswith("[") else LIGHT_GRAY
                if "STUDYSYNC" in value.upper() or "PROCESS" in value.upper():
                    fill = BLUE
                    set_cell_text(
                        cell,
                        value,
                        bold=True,
                        color=WHITE,
                        size=9,
                        align=WD_ALIGN_PARAGRAPH.CENTER,
                    )
                else:
                    set_cell_text(
                        cell,
                        value,
                        bold=value.startswith("["),
                        size=9,
                        align=WD_ALIGN_PARAGRAPH.CENTER,
                    )
                set_cell_shading(cell, fill)
            else:
                set_cell_text(cell, "", size=9)
                cell._tc.get_or_add_tcPr().append(OxmlElement("w:noWrap"))

    if column_widths:
        for row in table.rows:
            for index, width in enumerate(column_widths):
                row.cells[index].width = Cm(width)
    doc.add_paragraph()
    add_caption(doc, title)
    return table


def add_page_number(paragraph):
    add_text(paragraph, "Page ", size=9)
    run = paragraph.add_run()
    set_run_font(run, size=9)
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


def add_toc_field(paragraph):
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Table of Contents"
    fld_char_3 = OxmlElement("w:fldChar")
    fld_char_3.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)
    run._r.append(text)
    run._r.append(fld_char_3)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.54)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, before, after in [
        ("Heading 1", 14, 14, 8),
        ("Heading 2", 12, 12, 6),
        ("Heading 3", 11, 10, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ["List Bullet", "List Number"]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(11)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(header, "StudySync Student Course Registration System", size=9, color=DARK_GRAY)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer)


def add_title_page(doc):
    for _ in range(3):
        doc.add_paragraph()

    lines = [
        ("FEDERAL UNIVERSITY OF TECHNOLOGY, OWERRI", 14, True),
        ("(FUTO)", 13, True),
        ("", 12, False),
        ("SCHOOL OF INFORMATION AND COMMUNICATION TECHNOLOGY", 12, True),
        ("DEPARTMENT OF SOFTWARE ENGINEERING", 12, True),
        ("", 12, False),
        ("SOFTWARE ANALYSIS AND DESIGN REPORT", 16, True),
        ("ON", 12, True),
        ("STUDYSYNC STUDENT COURSE REGISTRATION SYSTEM", 16, True),
        ("", 12, False),
        ("A GROUP PROJECT REPORT", 12, True),
        ("", 12, False),
        (
            "Submitted in Partial Fulfilment of the Requirements for the Award "
            "of the Bachelor of Technology (B.Tech.) Degree in Software Engineering",
            11,
            False,
        ),
        ("", 12, False),
        ("Course: Software Analysis and Design", 11, False),
        ("Course Code: INS 204", 11, False),
        ("", 12, False),
        ("Submitted By", 11, True),
        ("[Insert Student Name(s) and Registration Number(s)]", 11, False),
        ("", 12, False),
        ("29th July, 2026", 11, False),
    ]

    for text, size, bold in lines:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(5)
        add_text(paragraph, text, size=size, bold=bold)

    doc.add_page_break()


def add_table_of_contents(doc):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(paragraph, "TABLE OF CONTENTS", size=14, bold=True, color=BLUE)
    doc.add_paragraph()
    toc = doc.add_paragraph()
    add_toc_field(toc)
    doc.add_paragraph()
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(note, "The table of contents updates automatically when opened in Microsoft Word.", size=9, italic=True, color=DARK_GRAY)
    doc.add_page_break()


def set_diagram_box(cell, text, dark=False):
    set_cell_text(
        cell,
        text,
        bold=True,
        color=WHITE if dark else None,
        size=8.2,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    set_cell_shading(cell, BLUE if dark else LIGHT_BLUE)
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "10")
        node.set(qn("w:color"), BLUE)
        borders.append(node)


def add_clean_diagram(doc, title, rows, widths=None):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_index, values in enumerate(rows):
        for column_index, raw_value in enumerate(values):
            value = str(raw_value)
            cell = table.rows[row_index].cells[column_index]
            if value.startswith("[") and value.endswith("]"):
                content = value[1:-1]
                set_diagram_box(cell, content, dark="STUDYSYNC" in content or "EXPRESS API" in content)
            else:
                set_cell_text(cell, value, bold=bool(value), size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
            if widths:
                cell.width = Cm(widths[column_index])
    doc.add_paragraph()
    add_caption(doc, title)
    return table


def add_manual_contents(doc):
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(heading, "TABLE OF CONTENTS", size=14, bold=True, color=BLUE)
    add_table(
        doc,
        ["Section", "Page"],
        [
            ["Title Page", "1"],
            ["Declaration and Certification", "2"],
            ["Acknowledgement and Abstract", "3"],
            ["Table of Contents", "4"],
            ["Chapter One: Introduction", "5"],
            ["Chapter Two: System Analysis", "8"],
            ["Chapter Three: System Design", "11"],
            ["Chapter Four: Implementation", "16"],
            ["Chapter Five: Testing, Conclusion and Recommendations", "19"],
            ["References and Lucidchart Guide", "22"],
        ],
        widths=[14.0, 2.0],
        font_size=10,
    )
    add_heading(doc, "List of Figures", 2)
    for item in [
        "Figure 2.1: StudySync context diagram",
        "Figure 3.1: Three-tier deployment architecture",
        "Figure 3.2: Role-based page architecture",
        "Figure 3.3: Level-zero process model",
        "Figure 3.4: Core database ERD",
    ]:
        add_bullet(doc, item)
    doc.add_page_break()


def preliminary_pages(doc):
    add_heading(doc, "DECLARATION", 1)
    add_body(doc, "We declare that this report presents the analysis, design, implementation, and testing of StudySync, a Student Course Registration System for a university tutorial centre. The work is prepared for academic assessment and has not been submitted elsewhere for the same award.")
    doc.add_paragraph()
    add_table(
        doc,
        ["Candidate", "Registration Number", "Signature and Date"],
        [["[Insert Student Name]", "[Insert Registration Number]", "____________________"]],
        widths=[6.0, 5.0, 5.0],
        font_size=10,
    )
    add_heading(doc, "CERTIFICATION", 1)
    add_body(doc, "This is to certify that the StudySync project report was carried out under approved academic supervision and satisfies the requirements of the Software Analysis and Design course, pending any corrections required by the assessment panel.")
    add_table(
        doc,
        ["Role", "Name", "Signature and Date"],
        [
            ["Project Supervisor", "____________________", "____________________"],
            ["Course Instructor", "____________________", "____________________"],
        ],
        widths=[5.0, 5.5, 5.5],
        font_size=10,
    )
    doc.add_page_break()

    add_heading(doc, "ACKNOWLEDGEMENT", 1)
    add_body(doc, "We acknowledge the guidance of our academic staff, project supervisor, classmates, and the staff whose comments helped to shape the system requirements. We also appreciate the open-source communities behind Node.js, Express, PostgreSQL, and the browser technologies used in the implementation.")
    add_heading(doc, "ABSTRACT", 1)
    add_body(doc, "StudySync is a web-based Student Course Registration System designed for a university tutorial centre. The system replaces fragmented paper and spreadsheet records with one role-based application for user accounts, student profiles, courses, semesters, enrolments, instructor profiles, instructor assignments, class lists, and management reports. Student profiles include faculty and an optional department. Every course is identified by a unique course code and course title.")
    add_body(doc, "The implemented solution uses separate HTML pages for each role, shared CSS, and small vanilla JavaScript modules on the frontend. A Node.js and Express REST API validates requests, applies business rules, and authenticates users with JSON Web Tokens. PostgreSQL is hosted through Supabase, the API is prepared for Render, and the static frontend is prepared for Vercel. The design prevents duplicate student-course-semester enrolments and limits each course to one active instructor assignment in a semester. The intentionally small scope makes the project practical to explain, test, maintain, and defend.")
    add_body(doc, "Keywords: student registration, course enrolment, tutorial centre, role-based access, Express, PostgreSQL, Supabase.")
    doc.add_page_break()


def chapter_one(doc):
    add_heading(doc, "CHAPTER ONE", 1)
    add_heading(doc, "INTRODUCTION", 1)
    add_heading(doc, "1.1 Background of the Study", 2)
    add_body(doc, "University tutorial centres provide additional academic support to students who want more practice in selected courses. A centre must know who its students are, the faculty and department of each student, the courses selected in a semester, and the instructor responsible for every tutorial class. When these records are kept in paper files or unrelated spreadsheets, staff spend time searching for information and may create duplicate or inconsistent records.")
    add_body(doc, "StudySync centralises these activities in a small web application. It does not attempt to replace a university portal. It is limited to the operations of a private tutorial centre serving university students, and it therefore uses the terms courses, semesters, instructors, and tutorials.")
    add_heading(doc, "1.2 Statement of the Problem", 2)
    for item in [
        "Student records may be incomplete, duplicated, or difficult to retrieve.",
        "A student may be entered twice for the same course in the same semester.",
        "Staff may not know the active instructor assigned to a course.",
        "Managers cannot quickly obtain course enrolment totals or current class lists.",
        "Manual records provide weak accountability for important changes.",
    ]:
        add_bullet(doc, item)
    doc.add_page_break()

    add_heading(doc, "1.3 Aim and Objectives", 2)
    add_body(doc, "The aim is to design and implement a beginner-friendly Student Course Registration System for a university tutorial centre.")
    add_heading(doc, "Specific Objectives", 3)
    for item in [
        "Provide secure login and role-based access for administrators, registration officers, managers, and instructors.",
        "Record student registration number, name, contact information, faculty, and optional department.",
        "Maintain courses using a unique course code and course title.",
        "Open and close semesters using an academic session and date range.",
        "Enrol active students in one or more active courses for an open semester.",
        "Assign one active instructor to a course in a semester.",
        "Produce enrolment totals, assignment reports, audit records, and instructor class lists.",
    ]:
        add_numbered(doc, item)
    add_heading(doc, "1.4 Scope of the System", 2)
    add_body(doc, "The first version includes account management, student records, course and semester setup, course enrolment, instructor profiles, instructor assignment, dashboards, reports, and class lists. It excludes payment, attendance, results, timetable generation, messaging, online tutorials, and payroll.")
    doc.add_page_break()

    add_heading(doc, "1.5 Significance of the Study", 2)
    add_table(
        doc,
        ["Stakeholder", "Expected Benefit"],
        [
            ["Registration Officer", "Faster registration, search, update, and enrolment of students."],
            ["Tutorial Centre Manager", "Clear course setup, semester control, instructor allocation, and reports."],
            ["Instructor", "A direct view of assigned courses and registered students."],
            ["Administrator", "Controlled user accounts, roles, status, and audit records."],
            ["Students", "More accurate registration records and fewer duplicate entries."],
        ],
        widths=[4.5, 11.5],
        font_size=9.5,
    )
    add_heading(doc, "1.6 Feasibility", 2)
    add_body(doc, "Technical feasibility is supported by familiar web technologies and hosted services. Economic feasibility is improved by free or low-cost tiers on Vercel, Render, and Supabase. Operational feasibility is supported by simple forms, tables, separate role pages, and loading feedback on every API action button.")
    add_heading(doc, "1.7 Method and Constraints", 2)
    add_body(doc, "The project follows a simplified waterfall sequence: problem identification, requirements analysis, database and interface design, implementation, and testing. Constraints include limited project time, dependence on internet connectivity for hosted deployment, and the cold-start behaviour that may occur on a free Render service.")


def chapter_two(doc):
    doc.add_page_break()
    add_heading(doc, "CHAPTER TWO", 1)
    add_heading(doc, "SYSTEM ANALYSIS", 1)
    add_heading(doc, "2.1 Analysis of the Existing Process", 2)
    add_body(doc, "The assumed existing process uses paper forms, notebooks, messaging applications, or separate spreadsheet files. Registration staff collect student details and selected courses, while a manager records instructor allocations separately. This arrangement makes cross-checking difficult because there is no shared identifier or automatic rule connecting a student, course, and semester.")
    add_heading(doc, "2.2 Fact-Finding Methods", 2)
    add_table(
        doc,
        ["Method", "Purpose", "Likely Finding"],
        [
            ["Interview", "Ask staff how registration and allocation are performed.", "Roles and approval responsibilities."],
            ["Observation", "Watch a registration from start to finish.", "Repeated writing and slow searches."],
            ["Document review", "Inspect forms, class lists, and spreadsheets.", "Required fields and duplicated data."],
            ["Code inspection", "Compare the design with the implemented repository.", "Actual routes, tables, pages, and constraints."],
        ],
        widths=[3.0, 6.0, 7.0],
        font_size=9,
    )
    add_heading(doc, "2.3 Proposed System", 2)
    add_body(doc, "The proposed system stores records in one PostgreSQL database and exposes controlled operations through an Express API. Users sign in and are redirected to real HTML pages for their role. The registration officer manages students and enrolments, the manager handles academic setup and instructor allocation, the administrator manages accounts, and instructors view only their own assigned classes.")
    doc.add_page_break()

    add_heading(doc, "2.4 Functional Requirements", 2)
    add_table(
        doc,
        ["ID", "Requirement", "Primary Role"],
        [
            ["FR-01", "Authenticate an active user and return a signed JWT.", "All roles"],
            ["FR-02", "Create, list, activate, or deactivate login accounts.", "Administrator"],
            ["FR-03", "Create, search, and update student profiles.", "Registration Officer"],
            ["FR-04", "Create and update courses with course code and course title.", "Manager"],
            ["FR-05", "Create and update semesters and their open or closed status.", "Manager"],
            ["FR-06", "Enrol an active student in selected active courses.", "Registration Officer"],
            ["FR-07", "Create instructor profiles and link instructor accounts.", "Manager"],
            ["FR-08", "Assign or cancel an instructor-course-semester allocation.", "Manager"],
            ["FR-09", "Display reports, dashboards, audit records, and class lists.", "Authorised roles"],
        ],
        widths=[2.0, 10.5, 3.5],
        font_size=8.7,
    )
    add_heading(doc, "2.5 Main Use Cases", 2)
    add_body(doc, "The major use cases are Manage Users, Manage Students, Manage Courses, Manage Semesters, Enrol Student, Manage Instructors, Assign Instructor, View Reports, and View My Classes. Each use case maps directly to an Express route and a visible HTML page where appropriate.")
    doc.add_page_break()

    add_heading(doc, "2.6 Non-Functional Requirements", 2)
    for item in [
        "Security: passwords are hashed, JWTs are verified, roles are checked, and database secrets remain on the server.",
        "Usability: pages use familiar forms and tables, readable labels, responsive layout, and visible loading states.",
        "Reliability: database constraints prevent duplicate enrolments and double active assignments.",
        "Performance: lists are ordered, report queries are filtered by semester, and the API compresses responses.",
        "Maintainability: frontend, backend, and database files have separate folders and small modules.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "2.7 Core Business Rules", 2)
    add_body(doc, "Only active students, courses, instructors, and accounts may participate in new work. Only open semesters accept enrolments and assignments. A student-course-semester combination is unique. Only one active instructor assignment is allowed for a course and semester. Cancelled records remain in the database to preserve history.")
    add_clean_diagram(
        doc,
        "Figure 2.1: StudySync context diagram",
        [
            ["[ADMINISTRATOR\naccounts and audit]", "<---->", "[STUDYSYNC SYSTEM\nrole-based registration and assignment]"],
            ["[REGISTRATION OFFICER\nstudents and enrolments]", "<---->", "[STUDYSYNC SYSTEM\nvalidated requests and responses]"],
            ["[TUTORIAL CENTRE MANAGER\nsetup, allocation and reports]", "<---->", "[STUDYSYNC SYSTEM\ncentral PostgreSQL records]"],
            ["[INSTRUCTOR\nassigned courses and class lists]", "<---->", "[STUDYSYNC SYSTEM\nrestricted instructor view]"],
        ],
        widths=[5.4, 2.0, 8.6],
    )

def chapter_three(doc):
    doc.add_page_break()
    add_heading(doc, "CHAPTER THREE", 1)
    add_heading(doc, "SYSTEM DESIGN", 1)
    add_heading(doc, "3.1 Design Approach", 2)
    add_body(doc, "StudySync uses a simple three-tier design. The presentation tier contains separate HTML pages and small JavaScript files. The application tier contains the Express routes, middleware, controllers, and validation helpers. The data tier contains the Supabase-hosted PostgreSQL tables, keys, checks, indexes, and triggers. This separation gives each technology one clear responsibility.")
    add_clean_diagram(
        doc,
        "Figure 3.1: Three-tier deployment architecture",
        [
            ["[PRESENTATION TIER\nSeparate HTML pages, CSS and vanilla JavaScript\nHosted on Vercel]"],
            ["|\nHTTPS requests with Authorization: Bearer JWT\nv"],
            ["[EXPRESS API ON RENDER\nRoutes -> middleware -> controllers -> validation]"],
            ["|\nSupabase JavaScript client with backend-only secret key\nv"],
            ["[DATA TIER\nPostgreSQL tables and constraints on Supabase]"],
        ],
        widths=[16.0],
    )
    add_heading(doc, "3.2 Architectural Decision", 2)
    add_body(doc, "The browser never receives the Supabase secret key and never queries the application tables directly. Row Level Security is enabled, while approved application operations pass through the Express API. This allows the backend to enforce validation, roles, business rules, and audit logging before database access.")
    doc.add_page_break()

    add_heading(doc, "3.3 Role-Based Page Design", 2)
    add_body(doc, "The frontend does not construct dashboards by injecting an entire page with JavaScript. Each role has its own folder and real HTML documents. JavaScript is limited to API communication, form submission, table population, status messages, and access checks.")
    add_clean_diagram(
        doc,
        "Figure 3.2: Role-based page architecture",
        [
            ["[ADMINISTRATOR]", "---->", "[client/admin/\ndashboard.html\nusers.html]", "---->", "[/api/users\n/api/reports/audit-logs]"],
            ["[REGISTRATION OFFICER]", "---->", "[client/registration/\ndashboard.html\nstudents.html\nenrolments.html]", "---->", "[/api/students\n/api/enrolments]"],
            ["[TUTORIAL CENTRE MANAGER]", "---->", "[client/manager/\ndashboard, courses, semesters,\ninstructors, assignments, reports]", "---->", "[/api/courses\n/api/semesters\n/api/instructors\n/api/assignments\n/api/reports]"],
            ["[INSTRUCTOR]", "---->", "[client/instructor/\ndashboard.html]", "---->", "[/api/dashboard/instructor\n/api/instructors/me/courses]"],
        ],
        widths=[3.5, 1.2, 6.0, 1.2, 4.1],
    )
    add_heading(doc, "3.4 Navigation Rule", 2)
    add_body(doc, "After login, the role in the JWT determines the correct home page. Browser routing improves usability, but Express role middleware remains the authoritative permission control.")
    doc.add_page_break()

    add_heading(doc, "3.5 Data Flow Design", 2)
    add_body(doc, "The level-zero process model below keeps actors, processes, and data stores in separate cells. Each connector occupies its own gutter, so arrows do not cross labels or record fields.")
    add_clean_diagram(
        doc,
        "Figure 3.3: Level-zero process model",
        [
            ["[ADMINISTRATOR]", "<---->", "[1.0 MANAGE ACCOUNTS]", "<---->", "[USERS / AUDIT_LOGS]"],
            ["[REGISTRATION OFFICER]", "<---->", "[2.0 MANAGE STUDENTS]", "<---->", "[STUDENTS]"],
            ["[REGISTRATION OFFICER]", "<---->", "[3.0 ENROL STUDENT]", "<---->", "[ENROLMENTS]"],
            ["[TUTORIAL CENTRE MANAGER]", "<---->", "[4.0 SET UP COURSES AND SEMESTERS]", "<---->", "[COURSES / SEMESTERS]"],
            ["[TUTORIAL CENTRE MANAGER]", "<---->", "[5.0 ASSIGN INSTRUCTORS]", "<---->", "[INSTRUCTORS / ASSIGNMENTS]"],
            ["[INSTRUCTOR]", "<---->", "[6.0 VIEW MY CLASSES]", "<---->", "[ASSIGNMENTS / ENROLMENTS]"],
        ],
        widths=[3.7, 1.2, 5.2, 1.2, 4.7],
    )
    add_heading(doc, "3.6 Main Enrolment Flow", 2)
    add_body(doc, "The officer selects an active student, an open semester, and one or more active courses. The API validates all referenced records, removes repeated course IDs, and upserts the rows using the student_id, course_id, semester_id conflict key. The database then returns the saved enrolments for display.")
    doc.add_page_break()

    add_heading(doc, "3.7 Entity Relationship Design", 2)
    add_body(doc, "The core ERD uses two associative entities. ENROLMENT resolves the many-to-many relationship between STUDENT and COURSE for a SEMESTER. INSTRUCTOR_ASSIGNMENT records the instructor responsible for a COURSE in a SEMESTER. COURSE and SEMESTER are shown twice only to keep connector routes outside the boxes; each repeated box represents the same database entity.")
    add_clean_diagram(
        doc,
        "Figure 3.4: Core database ERD with connector gutters",
        [
            ["[STUDENT\nid PK\nregistration_number UK\nfirst_name, last_name\nfaculty, department\nphone, email, status]", "1 ----<", "[ENROLMENT\nid PK\nstudent_id FK\ncourse_id FK\nsemester_id FK\nenrolled_at, status\nUK: student+course+semester]", ">---- 1", "[COURSE (SHARED)\nid PK\ncourse_code UK\ncourse_title UK\ndescription, level, status]"],
            ["", "", "|\nmany-to-one\nv", "", ""],
            ["", "", "[SEMESTER (SHARED)\nid PK\nsemester_name\nacademic_session\nstart_date, end_date, status\nUK: name+session]", "", ""],
            ["[INSTRUCTOR\nid PK\nuser_id FK/UK\nstaff_number UK\nfirst_name, last_name\nphone, email, qualification, status]", "1 ----<", "[INSTRUCTOR_ASSIGNMENT\nid PK\ninstructor_id FK\ncourse_id FK\nsemester_id FK\nassigned_at, status]", ">---- 1", "[COURSE (SAME ENTITY)\ncourse_id references COURSE]"],
            ["", "", "|\nmany-to-one\nv", "", ""],
            ["", "", "[SEMESTER (SAME ENTITY)\nsemester_id references SEMESTER]", "", ""],
        ],
        widths=[4.0, 1.3, 5.3, 1.3, 4.1],
    )
    add_body(doc, "A USER may optionally link to one INSTRUCTOR profile through instructors.user_id. A USER may also create many AUDIT_LOG records. Foreign keys preserve referential integrity, while status values preserve historical records instead of deleting them.")
    doc.add_page_break()

    add_heading(doc, "3.8 Data Dictionary", 2)
    add_table(
        doc,
        ["Table", "Purpose", "Important Fields and Rules"],
        [
            ["users", "Login accounts and roles.", "email UK; password_hash; role; is_active; failed_login_attempts; locked_until."],
            ["students", "Tutorial-centre student profiles.", "registration_number UK; faculty required; department optional; contact details; status."],
            ["instructors", "Instructor profiles.", "user_id optional FK/UK; staff_number UK; email UK; qualification; status."],
            ["courses", "Courses offered by the centre.", "course_code UK; course_title UK; description; level; status."],
            ["semesters", "Academic registration periods.", "semester_name + academic_session UK; dates; open/closed status."],
            ["enrolments", "Student course selections.", "student_id FK; course_id FK; semester_id FK; unique combination; status."],
            ["instructor_assignments", "Instructor allocation to a course.", "instructor_id FK; course_id FK; semester_id FK; assigned_at; status."],
            ["audit_logs", "Accountability history.", "user_id FK; action; details; created_at."],
        ],
        widths=[3.2, 4.3, 8.5],
        font_size=8.3,
    )
    add_heading(doc, "3.9 Normalisation to Third Normal Form", 2)
    add_body(doc, "The design is in Third Normal Form because student, instructor, course, semester, and account facts are stored once in their own tables. Associative tables contain only relationship keys and relationship attributes. Non-key values depend on the key, the whole key, and no other non-key attribute.")
    add_heading(doc, "3.10 Security Design", 2)
    add_body(doc, "Passwords are hashed with bcryptjs. JWTs expire after a configured period. Five failed logins cause a temporary lock. Helmet, CORS, rate limiting, validation helpers, Row Level Security, backend-only secrets, role middleware, and audit logs provide layered protection.")


def chapter_four(doc):
    doc.add_page_break()
    add_heading(doc, "CHAPTER FOUR", 1)
    add_heading(doc, "SYSTEM IMPLEMENTATION", 1)
    add_heading(doc, "4.1 Implemented Technology Stack", 2)
    add_table(
        doc,
        ["Layer", "Technology", "Reason"],
        [
            ["Frontend", "HTML5, CSS3, vanilla JavaScript", "Separate visible pages, small learning curve, and no framework build step."],
            ["Backend", "Node.js 20+, Express 5", "Straightforward REST routes and middleware."],
            ["Authentication", "JWT and bcryptjs", "Simple cross-service login and password hashing."],
            ["Database", "PostgreSQL through Supabase", "Hosted relational database with constraints and SQL schema."],
            ["Frontend deployment", "Vercel", "Static hosting for the client directory."],
            ["Backend deployment", "Render", "Hosted Node.js service with environment variables and health checks."],
            ["Optional static preview", "Cloudflare Wrangler configuration", "The root wrangler.toml can serve the client assets when required."],
        ],
        widths=[3.2, 5.0, 7.8],
        font_size=9,
    )
    add_heading(doc, "4.2 Repository Structure", 2)
    add_body(doc, "The repository is divided into client/, server/, and database/. The client contains role folders and shared assets. The server contains routes, controllers, middleware, configuration, scripts, utilities, and tests. The database folder contains schema.sql and an identical initial migration.")
    add_clean_diagram(
        doc,
        "Figure 4.1: Deployment request path",
        [
            ["[USER BROWSER]", "---->", "[VERCEL STATIC FRONTEND]", "---->", "[RENDER EXPRESS API]", "---->", "[SUPABASE POSTGRESQL]"],
        ],
        widths=[2.6, 1.0, 4.0, 1.0, 4.0, 1.0, 3.0],
    )
    doc.add_page_break()

    add_heading(doc, "4.3 API and Module Implementation", 2)
    add_table(
        doc,
        ["Method and Route", "Purpose", "Main Controller"],
        [
            ["POST /api/auth/login", "Validate credentials and return JWT.", "authController"],
            ["GET/POST/PATCH /api/users", "Manage login accounts and status.", "userController"],
            ["GET/POST/PATCH /api/students", "Manage student profiles and search.", "studentController"],
            ["GET/POST/PATCH /api/courses", "Manage course code, title, level, and status.", "courseController"],
            ["GET/POST/PATCH /api/semesters", "Manage academic periods.", "semesterController"],
            ["GET/POST/PATCH /api/instructors", "Manage instructor profiles.", "instructorController"],
            ["GET/POST/PATCH /api/enrolments", "Create, list, and cancel enrolments.", "enrolmentController"],
            ["GET/POST/PATCH /api/assignments", "Create, list, and cancel assignments.", "assignmentController"],
            ["GET /api/dashboard/*", "Return role dashboard summaries.", "dashboardController"],
            ["GET /api/reports/*", "Return enrolment, assignment, and audit reports.", "reportController"],
        ],
        widths=[5.0, 7.0, 4.0],
        font_size=8.3,
    )
    add_heading(doc, "4.4 Example Request Flow", 2)
    add_body(doc, "When a registration officer submits students.html, students.js sends JSON to POST /api/students. Authentication middleware verifies the JWT, role middleware checks permission, studentController validates the registration number, faculty, name, gender, and contact details, Supabase inserts the record, an audit log is written, and the frontend refreshes the existing HTML table.")
    doc.add_page_break()

    add_heading(doc, "4.5 User Interface Implementation", 2)
    add_table(
        doc,
        ["Role Folder", "HTML Pages", "Main Tasks"],
        [
            ["client/admin", "dashboard.html, users.html", "View totals, manage accounts, inspect audit records."],
            ["client/registration", "dashboard.html, students.html, enrolments.html", "Register students and save course enrolments."],
            ["client/manager", "dashboard.html, courses.html, semesters.html, instructors.html, assignments.html, reports.html", "Configure the centre and review reports."],
            ["client/instructor", "dashboard.html", "View assigned courses and student class lists."],
        ],
        widths=[3.8, 6.2, 6.0],
        font_size=8.8,
    )
    add_heading(doc, "4.6 Loading and Feedback Behaviour", 2)
    add_body(doc, "The shared api.js and ui.js modules place the clicked API button in a loading state. The button is disabled, its width is preserved, a spinner is shown, and aria-busy is applied until the request finishes. This behaviour covers sign in, sign out, form submission, search, cancellation, and other button-triggered API actions. Success or error messages are displayed as toasts.")
    add_heading(doc, "4.7 Deployment Configuration", 2)
    for item in [
        "render.yaml defines the Node service, server root directory, health endpoint, and environment variables.",
        "vercel.json publishes client/ and adds security headers for the static frontend.",
        "server/.env.example documents Supabase, JWT, CORS, port, and initial administrator settings.",
        "database/schema.sql is run in the Supabase SQL Editor before the API starts.",
        "wrangler.toml provides an optional static-assets configuration without changing the primary Vercel deployment.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "4.8 Implementation Status", 2)
    add_body(doc, "The core workflows described in the analysis are present in the codebase. The implementation intentionally avoids fees, attendance, results, timetables, and messaging so that the system remains focused and defensible.")


def chapter_five(doc):
    doc.add_page_break()
    add_heading(doc, "CHAPTER FIVE", 1)
    add_heading(doc, "TESTING, CONCLUSION AND RECOMMENDATIONS", 1)
    add_heading(doc, "5.1 Testing Strategy", 2)
    add_body(doc, "Testing combines automated unit checks, JavaScript syntax checks, database inspection, and browser workflow checks. The aim is to confirm validation, permissions, business rules, loading feedback, and responsive role pages.")
    add_table(
        doc,
        ["Test", "Expected Result", "Status"],
        [
            ["Valid required text", "Value is trimmed and accepted.", "Passed"],
            ["Blank required text", "Validation error is raised.", "Passed"],
            ["Email normalisation", "Email is converted to lower case.", "Passed"],
            ["Short password", "Password validation rejects it.", "Passed"],
            ["Wrong role", "Role middleware returns forbidden response.", "Passed"],
            ["Server JavaScript syntax", "Every server source file parses.", "Passed"],
            ["Frontend JavaScript syntax", "Every client script parses.", "Passed"],
            ["Schema and migration", "Both SQL files are identical.", "Passed"],
            ["API action loading", "Clicked button shows spinner and is disabled.", "Implemented"],
        ],
        widths=[5.2, 8.2, 2.6],
        font_size=8.8,
    )
    add_heading(doc, "5.2 Acceptance Scenarios", 2)
    add_body(doc, "A complete demonstration should create role accounts, open a semester, create courses with codes and titles, register students with faculty details, create and link instructors, enrol students, assign instructors, view reports, and display an instructor class list.")
    doc.add_page_break()

    add_heading(doc, "5.3 Risk and Maintenance Plan", 2)
    add_table(
        doc,
        ["Risk", "Control"],
        [
            ["Secret key exposure", "Keep Supabase secret and JWT secret only in Render or server/.env."],
            ["Duplicate records", "Use controller checks plus PostgreSQL unique constraints."],
            ["Unauthorised action", "Verify JWT and role on every protected route."],
            ["Free-service cold start", "Allow initial loading feedback and consider a paid service if usage grows."],
            ["Data loss", "Use Supabase backups and export important records periodically."],
            ["Terminology drift", "Keep code, pages, schema, and report aligned to course, semester, instructor, and tutorial."],
        ],
        widths=[5.0, 11.0],
        font_size=9,
    )
    add_heading(doc, "5.4 Limitations", 2)
    add_body(doc, "The system depends on internet access in production and currently uses a simple JWT session rather than refresh tokens. It does not process fees, attendance, results, timetables, communications, or online tutorial delivery. Integration tests require a configured Supabase project and suitable test records.")
    add_heading(doc, "5.5 Maintenance", 2)
    add_body(doc, "Maintenance should include dependency review, database backups, audit-log checks, user deactivation when staff leave, validation of CORS origins after deployment changes, and rerunning the automated tests after every modification.")
    doc.add_page_break()

    add_heading(doc, "5.6 Summary of Findings", 2)
    add_body(doc, "The analysis showed that the centre needs a shared registration record, duplicate prevention, clear instructor allocation, and simple role separation. The final implementation meets those needs using separate role pages, an Express API, and a normalised PostgreSQL schema.")
    add_heading(doc, "5.7 Conclusion", 2)
    add_body(doc, "StudySync provides a focused Student Course Registration System for a university tutorial centre. It records students by faculty and optional department, identifies every course by course code and course title, organises work by semester, prevents duplicate enrolments, tracks instructor assignments, and provides class lists and reports. The implementation matches the system analysis and database design while remaining small enough for a student team to explain confidently.")
    add_heading(doc, "5.8 Recommendations", 2)
    for item in [
        "Deploy the frontend, backend, and database separately on Vercel, Render, and Supabase.",
        "Use fictional data during demonstrations and never expose environment secrets.",
        "Train each staff role only on the pages required for that role.",
        "Back up the database and review audit logs regularly.",
        "Add future modules one at a time only after the core workflow remains stable.",
    ]:
        add_numbered(doc, item)
    add_heading(doc, "5.9 Suggested Future Enhancements", 2)
    add_body(doc, "Possible later modules include attendance, fee tracking, tutorial timetables, result records, notifications, data export, and stronger automated integration testing. These are extensions, not requirements of the first version.")


def references_and_guide(doc):
    doc.add_page_break()
    add_heading(doc, "REFERENCES", 1)
    references = [
        "Node.js. Node.js Documentation. https://nodejs.org/docs/",
        "Express.js. Express Web Framework Documentation. https://expressjs.com/",
        "PostgreSQL Global Development Group. PostgreSQL Documentation. https://www.postgresql.org/docs/",
        "Supabase. Supabase Documentation. https://supabase.com/docs",
        "Vercel. Vercel Documentation. https://vercel.com/docs",
        "Render. Render Documentation. https://render.com/docs",
        "Mozilla Developer Network. HTML, CSS, and JavaScript Web Documentation. https://developer.mozilla.org/",
        "OWASP Foundation. OWASP Web Application Security Guidance. https://owasp.org/",
    ]
    for reference in references:
        add_body(doc, reference)
    add_heading(doc, "APPENDIX A: LUCIDCHART RECREATION GUIDE", 1)
    add_body(doc, "Use standard rectangles for actors and processes, database/entity shapes for data stores, and crow's-foot connectors for the ERD. Turn on grid and snap. Place each entity in its own box, then route every relationship with elbow connectors through the empty gutters between boxes. Connector lines must enter the edge of a box and must never pass across the entity name or field list.")
    add_table(
        doc,
        ["Diagram", "Lucidchart Layout"],
        [
            ["Context", "Actors in one column, StudySync boundary in a second column, bidirectional arrows in a narrow gutter."],
            ["Architecture", "Vercel frontend above Render API above Supabase PostgreSQL, using vertical arrows."],
            ["Role pages", "Role, HTML folder/pages, and API groups in three columns with horizontal connector gutters."],
            ["DFD", "Actor, numbered process, and data store in three columns; one row per process."],
            ["ERD", "STUDENT-ENROLMENT-COURSE and INSTRUCTOR-INSTRUCTOR_ASSIGNMENT-COURSE, with SEMESTER below the associative entities. Use elbow routes around boxes."],
        ],
        widths=[3.2, 12.8],
        font_size=8.8,
    )
    add_heading(doc, "APPENDIX B: DEFENCE SUMMARY", 1)
    add_body(doc, "StudySync is not a full university portal. It is a private university tutorial-centre system. Students belong to faculties and may have departments; staff members are instructors; teaching sessions are tutorials; academic periods are semesters; and registrations are made for courses identified by course code and course title.")


def build_report():
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "StudySync Student Course Registration System"
    doc.core_properties.course = "Software Analysis and Design Report"
    doc.core_properties.author = "StudySync Project Team"
    doc.core_properties.keywords = "StudySync, student course registration, software analysis and design"
    doc.core_properties.comments = "Beginner-friendly SAD report prepared for academic defence."

    add_title_page(doc)
    preliminary_pages(doc)
    add_manual_contents(doc)
    chapter_one(doc)
    chapter_two(doc)
    chapter_three(doc)
    chapter_four(doc)
    chapter_five(doc)
    references_and_guide(doc)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(f"Created: {OUTPUT_PATH}")


if __name__ == "__main__":
    build_report()
